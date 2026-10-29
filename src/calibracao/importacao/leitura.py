"""Leitura da planilha DQ05: cronograma (lista mestra) + uma aba por instrumento.

Layout da aba de instrumento (igual nas abas da DQ05):
- cabeçalho: tipo em M3, código em AH3, faixa em M4, resolução em AH4, fabricante em M5,
  data da compra em AH5, periodicidade (meses) em M6;
- pontos da linha 11 até a linha "Laudo": ordem (A), ponto (B), tolerância (I);
- cada calibração é um bloco de 16 colunas a partir de M; a data fica na linha 8 e a incerteza
  na coluna +8 do bloco. Do histórico só se aproveita a incerteza da calibração mais recente.
"""

import datetime
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from calibracao.importacao.interpretacao import (
    interpretar_tolerancia,
    nominal_do_texto,
    normalizar,
    texto_exibido,
)
from calibracao.regras.localizacao import padronizar_localizacao

ABA_CRONOGRAMA = "Cronograma de Calibração"
PRIMEIRA_LINHA_CRONOGRAMA = 8
COL_CODIGO, COL_EQUIPAMENTO, COL_LOCALIZACAO = 1, 3, 6  # A, C, F

LINHA_PRIMEIRO_PONTO = 11
LINHAS_LIDAS = 60
COL_ORDEM, COL_PONTO, COL_TOLERANCIA = 1, 2, 9  # A, B, I
COL_PRIMEIRO_BLOCO, LARGURA_BLOCO, DESLOCAMENTO_INCERTEZA = 13, 16, 8  # M, 16 colunas, +8
MAXIMO_BLOCOS = 40  # a DQ05 tem espaço para ~32 calibrações por aba
LINHA_DATA_CALIBRACAO = 8
VERMELHO = "FFFF0000"


@dataclass
class PontoLido:
    linha: int
    ordem_original: str | None
    secao: str | None
    tipo: str
    descricao: str
    valor_nominal: float | None = None
    unidade: str | None = None
    tolerancia_inferior: float | None = None
    tolerancia_superior: float | None = None
    valor_minimo: float | None = None
    valor_maximo: float | None = None
    classe_exigida: str | None = None
    criterio: str | None = None  # texto da coluna "Tolerância" como aparece na planilha
    incerteza: float | None = None
    incerteza_data: datetime.date | None = None  # só para o relatório de conferência


@dataclass
class InstrumentoLido:
    codigo: str
    tipo: str
    localizacao: str | None
    aba: str | None = None
    faixa: str | None = None
    resolucao: str | None = None
    fabricante: str | None = None
    data_compra: str | None = None
    periodicidade_meses: int | None = None
    pontos: list[PontoLido] = field(default_factory=list)


@dataclass
class Excluido:
    codigo: str
    motivo: str


@dataclass
class ResultadoLeitura:
    instrumentos: list[InstrumentoLido]
    excluidos: list[Excluido]
    avisos: list[str]


def _grade(aba, max_linha: int) -> dict[tuple[int, int], object]:
    """Células preenchidas de uma aba: (linha, coluna) → célula."""
    grade = {}
    for linha in aba.iter_rows(max_row=max_linha):
        for celula in linha:
            if celula.value not in (None, ""):
                grade[(celula.row, celula.column)] = celula
    return grade


def _exibido(grade, linha: int, coluna: int) -> str | None:
    celula = grade.get((linha, coluna))
    return texto_exibido(celula.value, celula.number_format) if celula else None


def _linha_desativada(celulas) -> bool:
    """Linha do cronograma tachada ou com fundo vermelho = instrumento desativado."""
    for celula in celulas:
        if celula.value in (None, ""):
            continue
        if celula.font is not None and celula.font.strike:
            return True
        cor = celula.fill.fgColor if celula.fill is not None and celula.fill.fill_type else None
        if cor is not None and cor.type == "rgb" and cor.rgb == VERMELHO:
            return True
    return False


def ler_cronograma(planilha) -> tuple[list[tuple[str, str | None, str | None]], list[Excluido]]:
    """Lista mestra: (código, equipamento, localização) dos instrumentos ativos do cronograma."""
    ativos, excluidos = [], []
    aba = planilha[ABA_CRONOGRAMA]
    for linha in aba.iter_rows(min_row=PRIMEIRA_LINHA_CRONOGRAMA, max_col=COL_LOCALIZACAO):
        codigo = texto_exibido(linha[COL_CODIGO - 1].value)
        if not codigo:
            continue
        if _linha_desativada(linha):
            excluidos.append(Excluido(codigo, "Cronograma: linha tachada/vermelha (desativado)"))
            continue
        equipamento = texto_exibido(linha[COL_EQUIPAMENTO - 1].value)
        # Padronizada (Laboratório - Reserva…, - Orientação indicativa): pedido do usuário.
        localizacao = padronizar_localizacao(texto_exibido(linha[COL_LOCALIZACAO - 1].value))
        ativos.append((codigo, equipamento, localizacao))
    return ativos, excluidos


def _incerteza_mais_recente(grade, linha: int) -> tuple[float | None, datetime.date | None]:
    mais_recente = (None, None)
    for bloco in range(MAXIMO_BLOCOS):
        coluna = COL_PRIMEIRO_BLOCO + bloco * LARGURA_BLOCO + DESLOCAMENTO_INCERTEZA
        data = grade.get((LINHA_DATA_CALIBRACAO, coluna))
        valor = grade.get((linha, coluna))
        if data is None or not isinstance(data.value, datetime.datetime) or valor is None:
            continue
        if isinstance(valor.value, bool) or not isinstance(valor.value, int | float):
            continue
        dia = data.value.date()
        if mais_recente[1] is None or dia >= mais_recente[1]:
            mais_recente = (float(valor.value), dia)
    return mais_recente


def ler_pontos(grade) -> list[PontoLido]:
    """Pontos de uma aba de instrumento, já com seção, tipo e incerteza."""
    fim = next(
        (
            linha
            for linha in range(LINHA_PRIMEIRO_PONTO, LINHAS_LIDAS)
            if normalizar(_exibido(grade, linha, COL_ORDEM)).startswith("LAUDO")
        ),
        LINHAS_LIDAS,
    )
    linhas = []
    for linha in range(LINHA_PRIMEIRO_PONTO, fim):
        descricao = _exibido(grade, linha, COL_PONTO)
        tolerancia = _exibido(grade, linha, COL_TOLERANCIA)
        if descricao is None and tolerancia is None:
            continue
        linhas.append((linha, _exibido(grade, linha, COL_ORDEM), descricao, tolerancia))

    # Título de seção = linha de texto sem tolerância que tem a ordem vazia ou que vem logo antes de
    # pontos com tolerância. No 2º caso a seção vale só para esses pontos: itens de verificação que
    # vierem depois ficam fora dela (ex.: válvula: "Fechamento" → 8 ± 0,5; depois "Limpeza").
    pontos, secao, secao_so_com_tolerancia = [], None, False
    for indice, (linha, ordem, descricao, tolerancia) in enumerate(linhas):
        celula_ponto = grade.get((linha, COL_PONTO))
        numerico = celula_ponto is not None and isinstance(celula_ponto.value, int | float)
        if tolerancia is None and not numerico:
            proxima_tem_tolerancia = indice + 1 < len(linhas) and linhas[indice + 1][3] is not None
            if ordem is None or proxima_tem_tolerancia:
                secao, secao_so_com_tolerancia = descricao, ordem is not None
                continue
            if secao_so_com_tolerancia:
                secao, secao_so_com_tolerancia = None, False
            ponto = PontoLido(linha, ordem, secao, "VERIFICACAO", descricao)
        else:
            descricao = descricao or "(sem descrição na planilha)"
            ponto = PontoLido(linha, ordem, secao, "NUMERICO", descricao, criterio=tolerancia)
            if tolerancia is not None:
                interpretada = interpretar_tolerancia(tolerancia, descricao)
                ponto.tipo = interpretada.tipo
                ponto.tolerancia_inferior = interpretada.tolerancia_inferior
                ponto.tolerancia_superior = interpretada.tolerancia_superior
                ponto.valor_minimo = interpretada.valor_minimo
                ponto.valor_maximo = interpretada.valor_maximo
                ponto.unidade = interpretada.unidade
                ponto.classe_exigida = interpretada.classe_exigida
            if ponto.tipo == "NUMERICO":
                ponto.valor_nominal = (
                    float(celula_ponto.value) if numerico else nominal_do_texto(descricao)
                )
        ponto.incerteza, ponto.incerteza_data = _incerteza_mais_recente(grade, linha)
        pontos.append(ponto)
    return pontos


def _periodicidade(texto: str | None) -> int | None:
    try:
        return int(float(texto.replace(",", "."))) if texto else None
    except ValueError:
        return None


def ler_planilha(caminho: Path) -> ResultadoLeitura:
    planilha = openpyxl.load_workbook(caminho, read_only=True)
    try:
        ativos, excluidos = ler_cronograma(planilha)
        avisos: list[str] = []

        abas = {}
        for aba in planilha.worksheets:
            if aba.title == ABA_CRONOGRAMA:
                continue
            grade = _grade(aba, LINHAS_LIDAS)
            codigo = _exibido(grade, 3, 34)  # AH3
            if not codigo:
                continue
            if codigo in abas:
                avisos.append(f"Código {codigo} aparece em mais de uma aba; usada a primeira.")
                continue
            abas[codigo] = (aba.title, grade)

        codigos_ativos = {codigo for codigo, _, _ in ativos}
        codigos_desativados = {item.codigo for item in excluidos}
        for codigo in sorted(set(abas) - codigos_ativos - codigos_desativados):
            excluidos.append(Excluido(codigo, "Aba sem linha no cronograma (desativado)"))

        instrumentos = []
        for codigo, equipamento, localizacao in ativos:
            if codigo not in abas:
                instrumentos.append(
                    InstrumentoLido(codigo, equipamento or "(sem tipo)", localizacao)
                )
                continue
            titulo, grade = abas[codigo]
            instrumentos.append(
                InstrumentoLido(
                    codigo=codigo,
                    tipo=_exibido(grade, 3, 13) or equipamento or "(sem tipo)",  # M3
                    localizacao=localizacao,
                    aba=titulo,
                    faixa=_exibido(grade, 4, 13),  # M4
                    resolucao=_exibido(grade, 4, 34),  # AH4
                    fabricante=_exibido(grade, 5, 13),  # M5
                    data_compra=_exibido(grade, 5, 34),  # AH5
                    periodicidade_meses=_periodicidade(_exibido(grade, 6, 13)),  # M6
                    pontos=ler_pontos(grade),
                )
            )
        return ResultadoLeitura(instrumentos, excluidos, avisos)
    finally:
        planilha.close()

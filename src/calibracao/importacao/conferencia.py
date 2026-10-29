"""Planilha de conferência: tudo o que foi importado, para o usuário comparar com a DQ05."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from calibracao.importacao.leitura import ResultadoLeitura

CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor="424242")

TIPOS_PONTO = {
    "NUMERICO": "Numérico (nominal ± tolerância)",
    "LIMITE": "Limite (mínimo/máximo)",
    "CLASSE": "Classe exigida",
    "CRITERIO": "Critério (texto)",
    "VERIFICACAO": "Verificação (OK)",
}


def _numero(valor: float | None) -> str | None:
    return None if valor is None else f"{valor:g}".replace(".", ",")


def _aba(planilha, titulo: str, colunas: list[str], linhas: list[list]) -> None:
    aba = planilha.create_sheet(titulo)
    aba.append(colunas)
    for celula in aba[1]:
        celula.font, celula.fill = CABECALHO, FUNDO_CABECALHO
    for linha in linhas:
        aba.append(linha)
    for indice, coluna in enumerate(colunas, start=1):
        maior = max([len(str(coluna))] + [len(str(linha[indice - 1] or "")) for linha in linhas])
        aba.column_dimensions[get_column_letter(indice)].width = min(maior + 2, 60)
    aba.freeze_panes = "B2"
    aba.auto_filter.ref = aba.dimensions


def gerar_conferencia(resultado: ResultadoLeitura, caminho: Path) -> None:
    planilha = openpyxl.Workbook()
    planilha.remove(planilha.active)

    _aba(
        planilha,
        "Instrumentos",
        ["Código", "Tipo", "Localização", "Faixa", "Resolução", "Fabricante", "Data da compra",
         "Periodicidade (meses)", "Plano", "Nº de pontos", "Aba de origem"],
        [
            [i.codigo, i.tipo, i.localizacao, i.faixa, i.resolucao, i.fabricante, i.data_compra,
             i.periodicidade_meses, "Sim" if i.pontos else "SEM PLANO", len(i.pontos), i.aba]
            for i in resultado.instrumentos
        ],
    )
    _aba(
        planilha,
        "Pontos",
        ["Código", "Ordem", "Ordem na planilha", "Seção", "Tipo", "Ponto / descrição",
         "Nominal", "Unidade", "Tol. inferior", "Tol. superior", "Mínimo", "Máximo",
         "Classe exigida", "Tolerância (como na planilha)", "Incerteza (U)",
         "U veio da calibração de", "Linha na aba"],
        [
            [i.codigo, ordem, p.ordem_original, p.secao, TIPOS_PONTO[p.tipo], p.descricao,
             _numero(p.valor_nominal), p.unidade, _numero(p.tolerancia_inferior),
             _numero(p.tolerancia_superior), _numero(p.valor_minimo), _numero(p.valor_maximo),
             p.classe_exigida, p.criterio, _numero(p.incerteza),
             p.incerteza_data.strftime("%d/%m/%Y") if p.incerteza_data else None, p.linha]
            for i in resultado.instrumentos
            for ordem, p in enumerate(i.pontos, start=1)
        ],
    )
    _aba(
        planilha,
        "Fora do programa",
        ["Código", "Motivo"],
        [[e.codigo, e.motivo] for e in sorted(resultado.excluidos, key=lambda e: e.codigo)],
    )
    _aba(planilha, "Avisos", ["Aviso"], [[a] for a in resultado.avisos] or [["(nenhum)"]])

    caminho.parent.mkdir(parents=True, exist_ok=True)
    planilha.save(caminho)

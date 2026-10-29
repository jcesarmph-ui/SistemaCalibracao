"""Importação de uma planilha FICTÍCIA no formato da DQ05 (a real nunca vai para o Git)."""

import datetime
import sqlite3

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from calibracao.importacao.__main__ import main, sha256
from calibracao.importacao.gravacao import gravar_banco
from calibracao.importacao.leitura import ler_planilha


def _cabecalho(aba, tipo, codigo):
    aba["M3"], aba["AH3"] = tipo, codigo
    aba["M4"], aba["AH4"] = "0-150", 0.01
    aba["M5"], aba["AH5"] = "Fabricante Fictício", datetime.datetime(2020, 5, 4)
    aba["M6"] = 12


def _pontos(aba, linhas):
    for numero, (ordem, ponto, tolerancia) in enumerate(linhas, start=11):
        aba.cell(numero, 1, ordem)
        aba.cell(numero, 2, ponto)
        aba.cell(numero, 9, tolerancia)
    aba.cell(11 + len(linhas), 1, "Laudo ")


@pytest.fixture
def planilha_ficticia(tmp_path):
    planilha = openpyxl.Workbook()
    cronograma = planilha.active
    cronograma.title = "Cronograma de Calibração"
    linhas = [
        ("FIC-001", "Paquímetro", "Sala A"),
        ("FIC-002", "Micrômetro", "Reserva Laboratório"),
        ("FIC-003", "Trena", "Sala B"),  # desativado: tachado com fundo vermelho
        ("FIC-004", "Nível", "Laboratório"),  # sem aba própria
    ]
    for numero, (codigo, equipamento, local) in enumerate(linhas, start=8):
        cronograma.cell(numero, 1, codigo)
        cronograma.cell(numero, 3, equipamento)
        cronograma.cell(numero, 6, local)
    for coluna in (1, 3, 6):
        celula = cronograma.cell(10, coluna)
        celula.font = Font(strike=True)
        celula.fill = PatternFill("solid", fgColor="FFFF0000")

    paquimetro = planilha.create_sheet("Paquímetro FIC-001")
    _cabecalho(paquimetro, "Paquímetro", "FIC-001")
    _pontos(paquimetro, [
        (1, "Face Externa", None),
        (2, 0, "± 0,05"),
        (3, 150, "± 0,05"),
        (4, "Teste funcional - ligado", "Executa função"),
        (5, "Escala legível", None),
    ])
    # duas calibrações: a mais recente (bloco 2) vale para a incerteza
    paquimetro["U8"] = datetime.datetime(2023, 1, 10)
    paquimetro["AK8"] = datetime.datetime(2024, 1, 10)
    paquimetro["U12"], paquimetro["AK12"] = 0.01, 0.02
    paquimetro["U13"], paquimetro["AK13"] = 0.01, "-"

    micrometro = planilha.create_sheet("Micrômetro FIC-002")
    _cabecalho(micrometro, "Micrômetro", "FIC-002")
    _pontos(micrometro, [
        (1, 0, 0.007),
        (2, "Planeza", "2µm máx"),
        (3, "Resistência", "Mín. 600 MPa"),
        (4, "Classificação (Classe)", 1),
        (5, "30° (Q1)", "± 0,25°"),
        (None, "Inspeção visual", None),
        (6, "Painel sem defeitos", None),
    ])
    micrometro["I11"].number_format = '"± "0.000'

    for titulo, codigo in (("Trena FIC-003", "FIC-003"), ("Trena FIC-099", "FIC-099")):
        aba = planilha.create_sheet(titulo)
        _cabecalho(aba, "Trena", codigo)
        _pontos(aba, [(1, 300, "± 0,5")])
    planilha.create_sheet("Programação da planilha")

    caminho = tmp_path / "dq05_ficticia.xlsx"
    planilha.save(caminho)
    return caminho


def _por_codigo(resultado):
    return {i.codigo: i for i in resultado.instrumentos}


def test_lista_mestra_e_o_cronograma(planilha_ficticia):
    resultado = ler_planilha(planilha_ficticia)
    assert [i.codigo for i in resultado.instrumentos] == ["FIC-001", "FIC-002", "FIC-004"]
    motivos = {e.codigo: e.motivo for e in resultado.excluidos}
    assert "tachada" in motivos["FIC-003"]
    assert "sem linha no cronograma" in motivos["FIC-099"]


def test_instrumento_sem_aba_entra_sem_plano(planilha_ficticia):
    nivel = _por_codigo(ler_planilha(planilha_ficticia))["FIC-004"]
    assert (nivel.tipo, nivel.localizacao, nivel.pontos) == ("Nível", "Laboratório", [])


def test_cadastro_vem_da_aba_e_localizacao_do_cronograma(planilha_ficticia):
    paq = _por_codigo(ler_planilha(planilha_ficticia))["FIC-001"]
    assert (paq.tipo, paq.localizacao, paq.faixa, paq.resolucao) == (
        "Paquímetro", "Sala A", "0-150", "0,01"
    )
    assert (paq.fabricante, paq.data_compra, paq.periodicidade_meses) == (
        "Fabricante Fictício", "2020-05-04", 12
    )


def test_pontos_do_paquimetro(planilha_ficticia):
    pontos = _por_codigo(ler_planilha(planilha_ficticia))["FIC-001"].pontos
    assert [p.descricao for p in pontos] == [
        "0", "150", "Teste funcional - ligado", "Escala legível"
    ]  # "Face Externa" é título de seção, não ponto
    zero = pontos[0]
    assert (zero.secao, zero.tipo, zero.valor_nominal) == ("Face Externa", "NUMERICO", 0)
    assert (zero.tolerancia_inferior, zero.tolerancia_superior) == (-0.05, 0.05)
    assert (zero.incerteza, zero.incerteza_data) == (0.02, datetime.date(2024, 1, 10))
    # na calibração mais recente o ponto 150 ficou com "-": vale a última incerteza numérica
    assert (pontos[1].incerteza, pontos[1].incerteza_data) == (0.01, datetime.date(2023, 1, 10))
    assert (pontos[2].tipo, pontos[2].criterio) == ("CRITERIO", "Executa função")
    escala = pontos[3]
    assert (escala.tipo, escala.incerteza) == ("VERIFICACAO", None)
    # "Face Externa" vale só para os pontos com tolerância logo abaixo dela
    assert escala.secao is None


def test_pontos_do_micrometro(planilha_ficticia):
    pontos = _por_codigo(ler_planilha(planilha_ficticia))["FIC-002"].pontos
    zero, planeza, resistencia, classe, angulo, painel = pontos
    assert (zero.criterio, zero.tolerancia_superior) == ("± 0,007", 0.007)
    assert (planeza.tipo, planeza.valor_maximo, planeza.unidade) == ("LIMITE", 2, "µm")
    assert (resistencia.valor_minimo, resistencia.unidade) == (600, "MPa")
    assert (classe.tipo, classe.classe_exigida) == ("CLASSE", "1")
    assert (angulo.valor_nominal, angulo.unidade) == (30, "°")
    assert (painel.tipo, painel.secao) == ("VERIFICACAO", "Inspeção visual")


def test_banco_gravado_sem_calibracoes(planilha_ficticia, tmp_path):
    banco = tmp_path / "calibracao.db"
    gravar_banco(ler_planilha(planilha_ficticia), banco, "dq05_ficticia.xlsx", "abc")
    conexao = sqlite3.connect(banco)
    assert conexao.execute("SELECT COUNT(*) FROM instrumentos").fetchone()[0] == 3
    assert conexao.execute("SELECT COUNT(*) FROM planos").fetchone()[0] == 2
    assert conexao.execute("SELECT COUNT(*) FROM pontos_plano").fetchone()[0] == 10
    sem_plano = conexao.execute(
        "SELECT plano_id FROM instrumentos WHERE codigo = 'FIC-004'"
    ).fetchone()[0]
    assert sem_plano is None
    assert conexao.execute("PRAGMA journal_mode").fetchone()[0] == "delete"
    # O programa começa sem nenhuma calibração (o histórico antigo fica na planilha).
    for tabela in ("calibracoes", "pontos_calibracao", "desativacoes"):
        assert conexao.execute(f"SELECT COUNT(*) FROM {tabela}").fetchone()[0] == 0
    conexao.close()


def test_nunca_sobrescreve_banco(planilha_ficticia, tmp_path):
    banco = tmp_path / "calibracao.db"
    banco.write_bytes(b"banco antigo")
    with pytest.raises(FileExistsError):
        gravar_banco(ler_planilha(planilha_ficticia), banco, "x", "y")
    assert banco.read_bytes() == b"banco antigo"


def test_comando_completo_mantem_planilha_intacta(planilha_ficticia, tmp_path, capsys):
    antes = sha256(planilha_ficticia)
    banco, conferencia = tmp_path / "saida" / "calibracao.db", tmp_path / "saida" / "conf.xlsx"
    opcoes = ["--banco", str(banco), "--conferencia", str(conferencia)]
    assert main([str(planilha_ficticia), *opcoes]) == 0
    assert sha256(planilha_ficticia) == antes
    assert "Planilha original intacta" in capsys.readouterr().out
    abas = openpyxl.load_workbook(conferencia).sheetnames
    assert abas == ["Instrumentos", "Pontos", "Fora do programa", "Avisos"]
    assert main([str(planilha_ficticia), "--banco", str(banco)]) == 1  # já existe
